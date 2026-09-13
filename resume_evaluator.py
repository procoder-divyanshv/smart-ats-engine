import os
import json
from pathlib import Path
import time
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel , Field
from pypdf import PdfReader
from docx import Document

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key :
    raise ValueError("GROQ_API_KEY is not fetchable.")

client = Groq(api_key = my_api_key)
model = "openai/gpt-oss-120b"
# job_description = """
# Role Overview
# The AI Engineer Consultant will design, develop, test, deploy, and support enterprise grade AI solutions for SP Global. This role will focus on Python based AI agent development, Model Context Protocol server integration, Retrieval Augmented Generation, LLM enabled applications, and secure AI workflow automation.
# The consultant will work closely with SP Global's Enterprise Architecture, Product, Data Engineering, Application Engineering, and Governance teams to deliver production ready AI capabilities that can be integrated into enterprise platforms and business workflows.
# This is a hands-on delivery role requiring strong engineering skills, rapid onboarding, independent execution, and the ability to produce maintainable, well documented solutions.
# Engagement Expectations
# The consultant is expected to quickly understand SP Global's technology environment, delivery processes, architecture standards, data governance requirements, and business context.
# The consultant will be responsible for delivering defined project outcomes, producing high quality code and technical documentation, supporting implementation activities, and providing knowledge transfer to internal SP Global teams to ensure long term maintainability.
# Key Responsibilities
# Design, build, test, and deploy Python based AI agents capable of executing complex, multi-step workflows across enterprise systems.
# Develop and integrate MCP servers to enable secure, standardized, and context aware interactions between AI agents, enterprise APIs, databases, internal tools, and knowledge repositories.
# Build LLM powered applications using agentic frameworks, tool use orchestration, prompt engineering, and workflow automation patterns.
# Develop custom Retrieval Augmented Generation using embeddings, vector databases, metadata filtering, document chunking, and semantic search.
# Integrate AI services with enterprise applications using Python, C#, .NET, REST APIs, GraphQL, and SQL based data sources .
# Implement secure tool calling patterns, role-based access controls, audit logging, human in the loop approvals, and traceability mechanisms.
# Build reusable AI components, agent tools, prompt templates, evaluation scripts, and orchestration workflows.
# Collaborate with AI Architects and Data Modelers to ensure AI solutions use accurate, governed, and explainable data.
# Support deployment, monitoring, troubleshooting, and performance tuning of AI applications in cloud or enterprise environments.
# Create clear technical documentation covering architecture, prompts, data flows, APIs, MCP interfaces, dependencies, and operational procedures.
# Conduct knowledge transfer sessions with SP Global internal engineering teams.
# Required Qualifications
# Bachelors or masters degree in computer science, Software Engineering, Artificial Intelligence, Data Science, or a related field.
# Strong hands-on experience developing production grade applications using Python .
# Working knowledge of C# and .NET for enterprise application integration.
# Strong SQL skills for querying, validating, and integrating structured enterprise data.
# Experience building AI agents, LLM applications, prompt workflows, and tool calling systems.
# Practical understanding of MCP server architecture and enterprise system integration patterns.
# Experience with LLMs such as OpenAI, Anthropic Claude, AWS Bedrock models, Azure OpenAI, Google Gemini, or similar platforms.
# Hands on experience with frameworks such as LangChain, LlamaIndex, Semantic Kernel, AutoGen, CrewAI , or similar agentic AI frameworks.
# Experience with RAG pipelines, vector databases, embeddings, semantic search, and document processing.
# Experience developing APIs using FastAPI, Flask, .NET Web API, or similar frameworks.
# Understanding of authentication, authorization, secrets management, logging, observability, and secure coding practices.
# Preferred Qualifications
# Experience in financial services, market intelligence, ratings or regulated enterprise environments.
# Experience with AWS, Azure, or Google Cloud.
# Familiarity with vector databases such as OpenSearch, FAISS, Pinecone, Weaviate, Milvus, or pgvector.
# Experience with AI evaluation, hallucination detection, prompt regression testing, and LLM observability.
# Key Skills
# Python
# C#
# .NET core
# LLMs
# LangChain
# REST APIs
# LlamaIndex
# Semantic Kernel
# SQL
# Expected Deliverables
# Production ready Python based AI agents.
# MCP server integrations.
# RAG pipelines and retrieval workflows.
# API integrations with enterprise systems.
# Prompt templates, agent tools, and evaluation scripts.
# Technical documentation and deployment guides.
# Knowledge transfer sessions for SP Global teams.
# Role: Search Engineer
# Industry Type: IT Services & Consulting
# Department: Engineering - Software & QA
# Employment Type: Full Time, Permanent
# Role Category: Software Development
# Education
# UG: Any Graduate
# PG: Any Postgraduate
# """

job_description= """Information Technology Intern (AI Engineering)  



What You’ll Do 

Support maintaining production systems for high availability and reliability  

Help build monitoring and observability dashboards and alerts for system health  

Contribute to automation solutions for routine operational tasks  

Assist in developing and maintaining automated tests for production systems  

Support production deployments and change management for minimal downtime  

 

What They’re Looking For 

Currently pursuing a Bachelor’s in Computer Science, Software Engineering, or a related field  

Understanding of cloud computing concepts, particularly Microsoft Azure  

Intermediate Python programming skills  

Strong problem-solving skills  

Strong written and verbal communication in English  

 

Preferred Skills (not mandatory) 

Exposure to AI/ML products such as Azure Machine Learning and Databricks  

Familiarity with monitoring, observability, and logging implementation  

Experience with source control (GitHub), pull request reviews, and static analysis (e.g., SonarQube)  

Interest in writing tests with Pytest  

 

What You’ll Gain 

Hands-on experience with production AI/ML systems from Day 1  

Mentorship from experienced SRE and AI engineers  

Exposure to modern MLOps, monitoring, and automation tools  

A chance to work at the intersection of AI, reliability, and engineering  """
 # """Role Overview
# We are seeking a Robotics & Computer Vision Engineer to develop autonomous navigation, perception, and control systems for autonomous robotic platforms. This role involves building real-time vision pipelines, sensor fusion algorithms, and low-level hardware-software integration for edge deployment.

# Key Responsibilities
# - Design, test, and deploy computer vision pipelines for real-time object detection, pose estimation, and tracking using OpenCV and YOLO frameworks.
# - Develop and configure autonomous control nodes, state machines, and communication bridges in ROS 2 / C++.
# - Implement sensor fusion algorithms (Extended Kalman Filters, IMU, LiDAR integration) for localization, alignment, and drift-free motion.
# - Build and run high-fidelity simulations in Gazebo and MATLAB/Simulink before hardware deployment.
# - Integrate perception stacks with embedded microcontrollers (STM32, Arduino) and edge compute modules (NVIDIA Jetson).
# - Troubleshoot control architectures (PID loops, trajectory planning) and communication protocols (UART, SPI, MQTT).

# Required Qualifications
# - Bachelor's or Master's degree in Robotics, Computer Science, Mechatronics, Electrical Engineering, or a related field.
# - Strong programming proficiency in C++ and Python.
# - Hands-on experience with ROS 2 (Robot Operating System) and Linux (Ubuntu).
# - Solid background in Computer Vision and Deep Learning (OpenCV, PyTorch, YOLO, Detectron2).
# - Practical experience with robotics simulation tools such as Gazebo or PX4 Autopilot.

# Preferred Qualifications
# - Experience with embedded microcontrollers (STM32, Arduino) and serial hardware interfacing.
# - Understanding of sensor aggregation, LiDAR, IMU filtering, and state estimation (EKF).
# - Experience with robotic path planning, obstacle avoidance algorithms (MPPI), and pose estimation.

# Key Skills
# C++
# Python
# ROS 2
# OpenCV
# PyTorch
# YOLO
# Gazebo
# Linux
# Sensor Fusion

# Employment Type: Full Time, Permanent
# Education: UG - Any Graduate (Robotics/CS preferred), PG - Any Postgraduate"""
# """ Role Overview
# We are looking for a Data Analyst & Cloud BI Specialist to analyze complex datasets, build automated reporting systems, and deliver actionable insights across cross-functional teams. This role focuses on data cleaning, exploratory data analysis, statistical modeling, and developing interactive business intelligence dashboards.

# Key Responsibilities
# - Perform end-to-end data extraction, data cleaning, preprocessing, and statistical analysis on large datasets.
# - Design, develop, and maintain interactive dashboards and reports using Power BI and Tableau.
# - Implement quality control metrics, anomaly detection mechanisms, and data validation pipelines to ensure data accuracy.
# - Build predictive models and perform trend/time-series analysis using Python libraries (Pandas, NumPy, Matplotlib, Seaborn).
# - Collaborate with engineering and product teams to integrate analytics pipelines with cloud storage and serverless computing workflows.
# - Present data-driven insights, KPI tracking, and root cause analysis findings to stakeholders.

# Required Qualifications
# - Bachelor's or Master's degree in Computer Science, Artificial Intelligence, Data Science, Statistics, Mathematics, or a related field.
# - Proficiency in Python for data analysis and data manipulation (Pandas, NumPy).
# - Strong data visualization expertise using tools like Power BI, Tableau, Matplotlib, or Seaborn.
# - Solid understanding of statistical analysis, data cleaning methodologies, and exploratory data analysis (EDA).
# - Experience with version control systems (Git/GitHub) and Linux environments.

# Preferred Qualifications
# - Practical knowledge of cloud services (AWS S3, AWS Lambda, AWS Glue, Amazon Kinesis).
# - Experience in time series analysis, anomaly detection, or predictive modeling.
# - Familiarity with quality metrics, automated reporting workflows, and cross-functional team collaboration.

# Key Skills
# Python
# Pandas
# NumPy
# Data Cleaning
# Power BI
# Tableau
# Statistical Analysis
# AWS
# Git

# Employment Type: Full Time, Permanent
# Education: UG - Any Graduate (CSE/Data Science preferred), PG - Any Postgraduate"""
class JobD(BaseModel):
    role : str
    required_skills : list[str]
    preferred_skills : list[str]=[]
    minimum_experience : float | None =None
    education_requirements : list[str]=[]
    responsibilities : list[str]=[]

job_schema = JobD.model_json_schema()


system_prompt = f"""You are an expert HR assistant.
Your job is to analyse the job descriptions and extract the structured information from there.
Return only the valid JSON matching the following schema : {job_schema}

IMPORTANT :
Do not return the same schema itself.
Do not return the fields like properties , type, or title.
Fill the schema with actual description extracted from the job description.
If minimum_experience is not mentioned return null.
If information for a list is missing return an empty list.
Do not invent information.

"""


user_prompt = f"""
Analyse the following job description :
{job_description}

"""

message_system = {
    "role" : "system",
    "content" : system_prompt
}

message_user = {
    "role": "user",
    "content": user_prompt

}
response_format = {
    "type" : "json_object"
}
messages = [message_system, message_user]

response = client.chat.completions.create(model = model , messages = messages, temperature = 0.0, response_format = response_format)
answer = response.choices[0].message.content

raw_json = answer
job_data = json.loads(raw_json)
job = JobD(**job_data)


print(job.minimum_experience)
print(job.education_requirements)



class MatchResult(BaseModel):
    score : float
    details : dict 

class Experience(BaseModel):
    company : str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | list[str] | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name : str | None = None
    email: str | None = None
    phone : str | None = None
    total_experience : float | None = None
    skills : list[str] =[]
    experience : list[Experience] =[]
    project : list[str] = []   
    certifications : list[str] =[]

resume_schema = Resume.model_json_schema()

def parse_resume(resume_text):
    system_prompt="""
You are a expert resume parser
Extract information from the resume based on its meaning, not only based on exact headings.
Different resumes may use different headings: For example
Experience
Work History
Internships
Employement
Professional Experience

 These may all contain the relevant experience. Skills may also appear in the skill section, work experience, internships and projects.
 Return only valid JSON matching schema :
 {resume_schema}

 Important rules:
 1. Do not invent information.
 2. If a value is not available return null.
 3. If a list has no information return an empty list.
 4. Include internships inside Experiences.
 5. Extract skills mentioned in the entire resume.
    """
    user_prompt = f"Parse the following resume : {resume_text}"

    message_system = {
        "role" : "system",
        "content" : system_prompt

    }
    message_user = {
        "role" : "user", 
        "content" : user_prompt
    }
    messages =[message_system,message_user]
    response_format  = {
        "type" : "json_object"
    }

    response = client.chat.completions.create(model = model, messages = messages, response_format= response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume_json = Resume(**data)

    return resume_json

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text =""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text :
            text += page_text + '\n'
    return text

def read_docs(file_path) :
    reader = Document(file_path)
    text =""
    for paragraph in reader.paragraphs:
        if paragraph.text.strip() :
            text += paragraph.text + "\n"

    for table in reader.tables :
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() :
                    text+= cell.text + "\n"


    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf" :
        return read_pdf(file_path)
    elif file_path.suffix.lower() in [".docx", ".doc"] :
        return read_docs(file_path)
   
    else :
        return None


def final_score(job, resume) :
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are a strict and highly accurate technical recruiter. Compare the job description with the provided resume.

    STRICT RULES:
    * Thoroughly scan the entire resume for required skills. Do not miss skills embedded in project descriptions or certifications (e.g., AWS, Python).
    * Treat synonyms as matches (e.g., "AWS Cloud" matches "AWS").
    * Academic projects, hackathons, and student clubs DO NOT count as "Enterprise Production Experience".
    * Only mark 'Experience requirement met: True' if the candidate has full-time, professional industry experience matching the JD.

    JOB DESCRIPTION: {job.model_dump_json(indent=2)}
    RESUME: {resume.model_dump_json(indent=2)}
    
    Return your response strictly as a JSON object matching this schema:
    {match_schema}
    """
    message ={
        "role" : "user",
        "content" : prompt
    }
    messages = [message]
    reponse_format = {
        "type" : "json_object"
    }


    response = client.chat.completions.create(model = model , messages = messages, response_format = response_format)
    ans = response.choices[0].message.content

    data= json.loads(ans)
    return MatchResult(**data)



# actually retreiving the resumes from the folder 
resume_folder = Path("resumes")
all_results = []
for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf",".docx",".doc"] :
        continue

    print(f"\n Processing file : {file_path.name}")
    resume_text= read_resume(file_path)
    parsed_resume = parse_resume(resume_text)
    time.sleep(5)
    result = final_score(job, parsed_resume)
    time.sleep(5)
    print(f"Score : {result.score}")
    all_results.append(
        {
            "name" : parsed_resume.name,
            "score" :result.score,
            "details" : result.details,
        }
    )
all_results.sort(
    key =lambda candidate : candidate['score'],
    reverse= True

)
top_2 = all_results[:2]
worst_2 = all_results[-2:]

print("Top 2 Candidates : \n")
for candidate in top_2:
    print(f"{candidate['name']} - {candidate['score']}% \n")
    print(candidate['details'])
    print("\n")


print("Lowest 2 Candidates : \n")
for candidate in worst_2:
    print(f"{candidate['name']} - {candidate['score']}% \n")
    print(candidate['details'])
    print("\n")

if __name__ == "__main__":
    main()