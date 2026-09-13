import os
import io
import json
from typing import List
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from docx import Document

# Import your existing models and client logic from resume_evaluator2
from resume_evaluator2 import client, model, JobD, Resume, MatchResult

app = FastAPI(title="AI Resume Evaluator")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def extract_file_text(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    elif ext == ".docx":
        doc = Document(io.BytesIO(content))
        for p in doc.paragraphs:
            if p.text.strip():
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
    return text

def parse_jd(jd_text: str) -> JobD:
    job_schema = JobD.model_json_schema()
    system_prompt = f"""
    You are an expert HR Data Engineer. Analyze the job description and extract structured information.
    Return strictly valid JSON matching this schema: {job_schema}
    """
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": f"Analyse the following job description:\n{jd_text}"}
        ],
        temperature=0.0,
        response_format={"type": "json_object"}
    )
    return JobD(**json.loads(response.choices[0].message.content))

def parse_resume_content(resume_text: str) -> Resume:
    resume_schema = Resume.model_json_schema()
    system_prompt = f"""
    You are an expert resume parser.
    Extract information from the resume based on its meaning, not only based on exact headings.
    Return only valid JSON matching schema: {resume_schema}
    """
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": f"Parse the following resume: {resume_text}"}
        ],
        response_format={"type": "json_object"}
    )
    return Resume(**json.loads(response.choices[0].message.content))

def compute_match(job: JobD, resume: Resume) -> MatchResult:
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are a highly accurate technical recruiter. Compare the job description with the provided resume.
    JOB DESCRIPTION: {job.model_dump_json(indent=2)}
    RESUME: {resume.model_dump_json(indent=2)}
    Return strictly JSON matching: {match_schema}
    """
    response = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"}
    )
    return MatchResult(**json.loads(response.choices[0].message.content))

@app.post("/api/evaluate")
async def evaluate_resumes(
    job_description: str = Form(...),
    files: List[UploadFile] = File(...)
):
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description is required.")
    if not files:
        raise HTTPException(status_code=400, detail="At least one resume is required.")

    parsed_job = parse_jd(job_description)
    results = []

    for file in files:
        if not file.filename.lower().endswith((".pdf", ".docx")):
            continue

        content = await file.read()
        raw_text = extract_file_text(file.filename, content)
        if not raw_text.strip():
            continue

        parsed_res = parse_resume_content(raw_text)
        score_res = compute_match(parsed_job, parsed_res)

        results.append({
            "filename": file.filename,
            "candidate_name": parsed_res.name or file.filename,
            "score": score_res.score,
            "details": score_res.details.model_dump()
        })

    results.sort(key=lambda x: x["score"], reverse=True)
    return results

# Mount static frontend directory
if os.path.exists("static"):
    app.mount("/", StaticFiles(directory="static", html=True), name="static")